"""Test-only independent readers; not product acquisition or algorithm integration."""
import csv
import json
from pathlib import Path
import numpy as np
from scipy.io import loadmat
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent / "artifacts"
checks = []
relative = [[1,4,1,0,5],[1,5,1.25,.25,-6],[2,0,0,0,10],[2,1,.5,.5,20],[2,2,1,1,30]]
source = [[1,4,1,1,5],[1,5,1.25,1.25,-6],[2,0,0,0,10],[2,1,.5,.5,20],[2,2,1,1,30]]
files = [(f"{base}.{ext}", relative) for base in ["signal-data/core-compare", "signal-data/api-compare", "desktop-signal/relative"] for ext in ["json","csv","txt","xlsx","mat"]]
files += [("desktop-signal/source.json",source), ("desktop-signal/segment.json",[[1,5],[1.25,-6]]), ("desktop-signal/channel.json",[[1,5],[1.25,-6]]), ("desktop-signal/restored.json",relative), ("desktop-signal/restored.mat",relative)]
for filename, expected in files:
    file = root / filename
    ext = file.suffix
    if ext == ".json":
        dataset = json.loads(file.read_text())
        actual, envelope = dataset["rows"], dataset
    elif ext in [".csv", ".txt"]:
        with file.open(encoding="utf-8-sig",newline="") as stream:
            lines = list(csv.reader(stream,delimiter="," if ext == ".csv" else "\t"))
        actual, envelope = [[float(v) for v in row] for row in lines[3:]], json.loads(lines[1][1])
    elif ext == ".xlsx":
        book = load_workbook(file,read_only=True,data_only=False)
        sheet = book["结果数据"]
        actual = [list(row) for row in sheet.iter_rows(min_row=2,values_only=True)]
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        envelope = json.loads("".join(row[2] for row in book["元信息"].iter_rows(values_only=True) if row[0] == "metadata_json"))
        book.close()
    else:
        mat = loadmat(file)
        actual = mat["values"].tolist()
        envelope = json.loads(bytes(mat["metadata_json_utf8"].flatten().astype(np.uint8)).decode("utf-8"))
        assert envelope["numericColumnIndices"] == list(range(len(expected[0])))
    np.testing.assert_allclose(actual,expected,rtol=0,atol=0)
    metadata = envelope["metadata"]
    assert envelope["columns"][-1]["unit"] == "mm/s"
    assert metadata["workflow"] == "signal-input-v1"
    assert metadata["source"] == "import"
    assert metadata["inputs"][0]["segment"]["startIndex"] == 4
    assert metadata["inputs"][0]["task"]["sampleRate"] == 4
    if len(expected[0]) == 5:
        assert metadata["inputs"][1]["task"]["sampleRate"] == 2
        assert metadata["inputs"][0]["sampleCount"] == 2
        assert metadata["inputs"][1]["sampleCount"] == 3
    checks.append({"file":filename,"status":"pass"})
(root / "signal-data/independent-readers.json").write_text(json.dumps(checks,ensure_ascii=False,indent=2))
print(f"PASS {len(checks)} actual signal exports independently read: source/display times, indices, units, samples and frozen input metadata")
