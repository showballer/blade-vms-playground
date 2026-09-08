"""Independent QA readers for actual product downloads. Not product Python integration."""
import csv
import json
import re
import unicodedata
from pathlib import Path
import numpy as np
from scipy.io import loadmat
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parent.parent / "artifacts/report-fidelity"
expected = json.loads((root / "expected.json").read_text())
checks = []
task = expected["task"]
wave = [[i / task["sampleRate"], value] for i, value in enumerate(task["samples"])]
for format in ["csv", "txt", "xlsx", "mat", "json"]:
    file = root / ("wave." + format)
    if format == "json":
        envelope = json.loads(file.read_text())
        actual = envelope["rows"]
    elif format in ["csv", "txt"]:
        with file.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter="," if format == "csv" else "\t"))
        envelope = json.loads(rows[1][1])
        actual = [[float(v) for v in row] for row in rows[3:]]
    elif format == "xlsx":
        book = load_workbook(file, read_only=True, data_only=False)
        sheet = book["结果数据"]
        actual = list(sheet.iter_rows(min_row=2, values_only=True))
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        envelope = json.loads("".join(row[2] for row in book["元信息"].iter_rows(values_only=True) if row[0] == "metadata_json"))
        book.close()
    else:
        mat = loadmat(file)
        actual = mat["values"].tolist()
        envelope = json.loads(bytes(mat["metadata_json_utf8"].flatten().astype(np.uint8)).decode("utf-8"))
        assert envelope["numericColumnIndices"] == [0, 1]
    np.testing.assert_allclose(actual, wave, rtol=0, atol=0)
    assert envelope["metadata"]["report"]["id"] == expected["reportId"]
    assert envelope["metadata"]["report"]["frozen"] is True
    assert envelope["metadata"]["task"]["sampleRate"] == 100
    assert envelope["metadata"]["task"]["note"] == "原始工况说明"
    assert envelope["columns"][1]["unit"] == "mm/s"
    checks.append({"file": file.name, "result": "pass"})

for file, rows in [("spectrum.json", expected["fft"]), ("deleted-source-wave.json", wave)] + [("modal-" + key + ".json", data["rows"]) for key, data in expected["modal"].items()]:
    data = json.loads((root / file).read_text())
    assert data["rows"] == rows
    assert data["metadata"]["report"]["id"] == expected["reportId"]
    checks.append({"file": file, "result": "pass"})

for file in ["saved.pdf", "word-pages/saved.pdf"]:
    pdf = PdfReader(root / file)
    # Embedded CJK fonts can map equivalent glyphs to compatibility radicals.
    # Read in geometric order: vector chart labels are painted after body text in Chromium.
    extracted = "".join(page.extract_text(extraction_mode="layout") for page in pdf.pages)
    text = re.sub(r"\s+", "", unicodedata.normalize("NFKC", extracted.translate(str.maketrans({"⻓":"长", "⻚":"页"}))))
    text = re.sub(r"第\d+页/共\d+页", "", text)
    for chapter in ["对象信息", "采集参数", "原始波形", "降噪与截取", "FFT频谱", "包络与时频", "模态参数", "损伤识别", "结论与建议"]:
        assert chapter in text, (file, chapter)
    assert text.count("长参数分页验收。") == 70, (file, "Long parameter truncated")
    assert "4096" in text and "8192" not in text
    assert "三阶弯曲" in text and "结果/位置:叶尖" in text
    assert "原始工况说明" in text and "后来修改的工况" not in text
    if file == "saved.pdf":
        assert "原始幅值" in text, "SVG waveform legend missing from the actual PDF"
    for page in pdf.pages:
        assert abs(float(page.mediabox.width) - 595.3) < 2
        assert abs(float(page.mediabox.height) - 841.9) < 2
    checks.append({"file": file, "pages": len(pdf.pages), "result": "pass", "scope": "A4 and complete extracted content; visual review recorded separately"})
(root / "independent-readers.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2))
print(f"PASS {len(checks)} actual report downloads independently read; frozen matrices, metadata, A4 and complete long content")
