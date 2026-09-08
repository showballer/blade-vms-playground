"""Independent readers for actual modal exports; QA only, not algorithm integration."""
import csv
import json
from datetime import datetime, timezone
from pathlib import Path
import numpy as np
from scipy.io import loadmat
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent / "artifacts"
checks = []


def read_export(filename):
    path = root / filename
    if path.suffix == ".json":
        document = json.loads(path.read_text(encoding="utf-8"))
        rows, reader = document["rows"], "json"
    elif path.suffix in (".csv", ".txt"):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            table = list(csv.reader(stream, delimiter="," if path.suffix == ".csv" else "\t"))
        document = json.loads(table[1][1])
        rows = [[float(cell) if document["columns"][i]["type"] == "number" else cell for i, cell in enumerate(row)] for row in table[3:]]
        reader = "csv.reader"
    elif path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        rows = [list(row) for row in book["结果数据"].iter_rows(min_row=2, values_only=True)]
        assert all(cell.data_type != "f" for row in book["结果数据"] for cell in row)
        chunks = [row for row in book["元信息"].iter_rows(values_only=True) if row[0] == "metadata_json"]
        assert [row[1] for row in chunks] == list(range(1, len(chunks) + 1))
        document = json.loads("".join(row[2] for row in chunks))
        reader = "openpyxl"
        book.close()
    else:
        assert path.suffix == ".mat"
        result = loadmat(path)
        document = json.loads(bytes(result["metadata_json_utf8"].flatten().astype(np.uint8)).decode("utf-8"))
        indices = result["numeric_column_indices"].flatten().astype(int).tolist()
        assert indices == document["numericColumnIndices"] and document["indexBase"] == 0
        rows = [[None] * len(document["columns"]) for _ in range(document["rowCount"])]
        for row_index, numeric_row in enumerate(result["values"].tolist()):
            for i, value in zip(indices, numeric_row):
                rows[row_index][i] = value
        for column in document["textColumns"]:
            for row_index, value in enumerate(column["values"]):
                rows[row_index][column["columnIndex"]] = value
        reader = "scipy.io.loadmat"
    assert document["schema"] == "blade-vms/dataset-v1"
    assert document.get("rowCount", len(rows)) == len(rows)
    return rows, document, reader


expected = {
    "parameters": [[1, .6, 1.2, "一阶弯曲"], [2, 1.8, 1.8, "二阶弯曲"], [3, 3.5, 2.1, "三阶弯曲"]],
    "points": [[.2, .4, .3], [.4, .7, .48], [.6, .6, .4], [.8, .9, .58], [1, .8, .51]],
    "shape": [[1, "P1 叶根", 0, .15, 0], [1, "P3 叶尖", 1, 1, 0], [2, "P1 叶根", 0, .2, 0], [2, "P3 叶尖", 1, .65, 0], [3, "P1 叶根", 0, .3, 0], [3, "P3 叶尖", 1, -.9, 0]],
}
files = [(f"modal-data/{prefix}{key}.{ext}", key, False) for prefix in ("", "api-") for key in ("parameters", "points", "shape", "baseline") for ext in ("json", "csv", "txt", "xlsx", "mat")]
files += [(f"modal-ui/{key}.json", key, True) for key in ("parameters", "points", "shape", "baseline")]
files += [("modal-ui/points.xlsx", "points", True), ("modal-ui/history-baseline.json", "baseline", True), ("modal-ui/history-shape.mat", "shape", True)]
for filename, key, ui in files:
    rows, document, reader = read_export(filename)
    if key == "baseline":
        baselines = [.62] if ui else [.61, 1.82, 3.55]
        wanted = [[i + 1, frequency, [.6, 1.8, 3.5][i], round(([.6, 1.8, 3.5][i] - frequency) / frequency * 100, 8)] for i, frequency in enumerate(baselines)]
    else:
        wanted = expected[key]
    assert rows == wanted, filename
    metadata = document["metadata"]
    assert metadata["workflow"] == "modal-preview-v1" and metadata["demo"] is True
    assert metadata["source"] == "demo" and metadata["datasetKey"] == key
    assert metadata["task"]["environment"] == "原始环境"
    assert metadata["parameters"]["points"] == ["P1", "P3"]
    assert metadata.get("analysis", {}).get("result", {}).get("datasets") is None
    assert "未根据所选任务" in metadata["note"]
    if key == "parameters":
        assert [c["unit"] for c in document["columns"]] == ["", "Hz", "%", ""]
    elif key == "points":
        assert [c["key"] for c in document["columns"]] == ["frequency", "P1", "P3"]
    checks.append({"file": filename, "reader": reader, "status": "pass"})

native_shape = [[1, "P2 叶中", .5, .55, 0], [2, "P2 叶中", .5, -.8, 0], [3, "P2 叶中", .5, .6, 0]]
for name, wanted in [
    ("points.json", [[.2, .36], [.4, .58], [.6, .49], [.8, .72], [1, .64]]),
    ("shape.mat", native_shape), ("restored-shape.xlsx", native_shape),
    ("restored-parameters.json", expected["parameters"]),
]:
    filename = "desktop-modal/" + name
    rows, document, reader = read_export(filename)
    assert rows == wanted
    assert document["metadata"]["parameters"]["points"] == ["P2"]
    assert document["metadata"]["task"]["environment"] == "原始环境"
    assert document["metadata"]["demo"] is True
    checks.append({"file": filename, "reader": reader, "status": "pass"})

(root / "modal-data/independent-readers.json").write_text(json.dumps({"time": datetime.now(timezone.utc).isoformat(), "checks": checks}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"PASS {len(checks)} actual modal exports: four result types, selected points, units, Chinese labels, baseline precision and frozen provenance")
