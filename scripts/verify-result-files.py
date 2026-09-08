"""Independent QA readers only; not application/Python-algorithm integration."""
import csv
import json
from pathlib import Path
import numpy as np
from scipy.io import loadmat
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent / "artifacts"
checks = []

for filename, expected in [
    ("result-data/raw.mat", [[0, 2], [.25, -4], [.5, 6], [.75, -8]]),
    ("result-data/api-cut.mat", [[.25, -4], [.5, 6]]),
    ("desktop-result-smoke/saved.mat", [[1.21], [3.4], [6.7], [3.1], [.02], [.6], [842], [28]]),
]:
    result = loadmat(root / filename)
    np.testing.assert_allclose(result["values"], expected, rtol=0, atol=0)
    metadata = json.loads(bytes(result["metadata_json_utf8"].flatten().astype(np.uint8)).decode("utf-8"))
    assert metadata["rowCount"] == len(expected)
    assert metadata["indexBase"] == 0
    if filename.endswith("saved.mat"):
        assert metadata["textColumns"][0]["values"][5] == "主频"
        assert metadata["textColumns"][1]["values"][5] == "Hz"
        assert metadata["metadata"]["task"]["environment"] == "导出验收"
    checks.append({"file": filename, "reader": "scipy.io.loadmat", "status": "pass"})

for suffix, delimiter in [("csv", ","), ("txt", "\t")]:
    filename = "result-data/api-cut." + suffix
    with (root / filename).open(encoding="utf-8-sig", newline="") as stream:
        records = list(csv.reader(stream, delimiter=delimiter))
    metadata = json.loads(records[1][1])
    assert metadata["metadata"]["analysis"]["parameters"]["sourceRange"] == {"startIndex": 1, "endExclusive": 3}
    assert [[float(cell) for cell in row] for row in records[3:]] == [[.25, -4], [.5, 6]]
    checks.append({"file": filename, "reader": "csv.reader", "status": "pass"})

for filename, first in [("result-data/api-cut.xlsx", [.25, -4]), ("desktop-result-smoke/features.xlsx", ["RMS", 1.21, "mm/s"]), ("desktop-result-smoke/raw.xlsx", [0, 99])]:
    book = load_workbook(root / filename, read_only=True, data_only=False)
    data = book["结果数据"]
    assert list(next(data.iter_rows(min_row=2, max_row=2, values_only=True))) == first
    chunks = [row[2] for row in book["元信息"].iter_rows(values_only=True) if row[0] == "metadata_json"]
    metadata = json.loads("".join(chunks))
    assert metadata["rowCount"] == data.max_row - 1
    assert all(cell.data_type != "f" for row in data for cell in row)
    book.close()
    checks.append({"file": filename, "reader": "openpyxl", "status": "pass"})

(root / "result-data/independent-readers.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"PASS {len(checks)} actual exported files independently read; numbers, Chinese metadata, units and snapshots preserved")
