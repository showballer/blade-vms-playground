"""Read actual rotation exports independently; QA only, not an algorithm service."""
import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat

root = Path(__file__).resolve().parent.parent / "artifacts"
checks = []
expected = [[1, 0, 0, 1], [1, 1, .25, -2], [3, 4, 1, 5], [3, 5, 1.25, -6]]


def read_export(filename):
    path = root / filename
    if path.suffix == ".json":
        document = json.loads(path.read_text(encoding="utf-8"))
        rows = document["rows"]
        reader = "json"
    elif path.suffix in (".csv", ".txt"):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            table = list(csv.reader(stream, delimiter="," if path.suffix == ".csv" else "\t"))
        assert table[1][0] == "# metadata_json"
        document = json.loads(table[1][1])
        rows = [[float(cell) for cell in row] for row in table[3:]]
        assert document["rowCount"] == len(rows)
        reader = "csv.reader"
    elif path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["结果数据"]
        rows = [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        chunks = [row for row in book["元信息"].iter_rows(values_only=True) if row[0] == "metadata_json"]
        assert [row[1] for row in chunks] == list(range(1, len(chunks) + 1))
        document = json.loads("".join(row[2] for row in chunks))
        assert document["rowCount"] == len(rows)
        book.close()
        reader = "openpyxl"
    else:
        assert path.suffix == ".mat"
        result = loadmat(path)
        document = json.loads(bytes(result["metadata_json_utf8"].flatten().astype(np.uint8)).decode("utf-8"))
        rows = result["values"].tolist()
        assert result["numeric_column_indices"].flatten().tolist() == [0, 1, 2, 3]
        assert document["indexBase"] == 0 and document["rowCount"] == len(rows)
        reader = "scipy.io.loadmat"
    assert document["schema"] == "blade-vms/dataset-v1"
    return rows, document, reader


files = [f"rotation-data/{name}.{ext}" for name in ("cut", "api-cut") for ext in ("csv", "txt", "xlsx", "mat", "json")]
files += [f"desktop-rotation/{name}" for name in ("selected.json", "selected.csv", "selected.mat", "compared.json", "restored.json", "restored.xlsx")]
for filename in files:
    rows, document, reader = read_export(filename)
    np.testing.assert_allclose(rows, expected, rtol=0, atol=0)
    assert [c["key"] for c in document["columns"]] == ["segment", "sampleIndex", "time", "value"]
    assert document["columns"][3]["unit"] == "mm/s"
    metadata = document["metadata"]
    assert metadata["workflow"] == "rotation-segments-v1"
    assert metadata["source"] == "import" and metadata["demo"] is False
    assert re.fullmatch(r"[0-9a-f]{64}", metadata["sourceRevision"])
    assert metadata["keptSamples"] == 4 and metadata["excludedSamples"] == 7
    assert len(metadata["segments"]) == 6 and metadata["segments"][-1]["partial"] is True
    assert [s["keep"] for s in metadata["segments"]] == [True, False, True, False, False, False]
    assert [(s["startIndex"], s["endIndex"]) for s in metadata["excludedRanges"]] == [(2, 4), (6, 8), (8, 10), (10, 11)]
    assert metadata["excludedRanges"][0]["quality"] == "异常"
    assert metadata["excludedRanges"][0]["note"].startswith("人工排除")
    assert metadata["task"]["sampleRate"] == 4 and metadata["task"]["channels"][0]["unit"] == "mm/s"
    if not filename.startswith("rotation-data/cut."):
        assert metadata["task"]["environment"] == "原始环境"
    checks.append({"file": filename, "reader": reader, "status": "pass"})

rows, document, reader = read_export("desktop-rotation/manual.json")
np.testing.assert_allclose(rows, [[1, 1, .25, -2], [1, 2, .5, 3], [1, 3, .75, -4], [1, 4, 1, 5]], rtol=0, atol=0)
assert document["metadata"]["parameters"]["mode"] == "manual"
assert len(document["metadata"]["segments"]) == 1 and document["metadata"]["excludedSamples"] == 0
checks.append({"file": "desktop-rotation/manual.json", "reader": reader, "status": "pass"})

rows, document, reader = read_export("desktop-rotation/order.json")
assert rows == [["1×", 1, 1, 1.8], ["fb", 2, 2, 1.25], ["2×", 2, 2, .8], ["3×", 3, 3, .5]]
assert document["metadata"]["demo"] is True and "未执行 FFT" in document["metadata"]["note"]
checks.append({"file": "desktop-rotation/order.json", "reader": reader, "status": "pass"})

(root / "rotation-data/independent-readers.json").write_text(json.dumps({"time": datetime.now(timezone.utc).isoformat(), "checks": checks}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"PASS {len(checks)} actual rotation exports: original indices/time/amplitudes, Chinese notes, excluded segments, frozen metadata and explicit demo order values")
